"""
app/services/part3_lining/service.py
====================================
Part 3 - Lining Designer service.

Input  : the values the user confirmed in the wizard (vessel profile from teammate 1's extraction,
         design basis, lining concept zones, nozzles) + the brick library.
Output : DXF drawings (lining GA with nozzles, brick shapes, layers/brick layout) with PNG/PDF
         previews, the course schedule, bottom ring table, volumes and checks.
"""
from __future__ import annotations

import csv
import datetime
import json
import math
import re
from pathlib import Path

from .lining.adapter import build_engine_inputs
from .lining.bom import build_bom
from .lining.checks import run_checks
from .lining.bricks import band_kind, bottom_rings, course_schedule, quantities
from .lining.drawings import BrickSheet, LayerSheet, LiningSheet, aci_hex, assign_mark_colors, mark_info, save_previews
from .lining.geometry import region_area_moment, reverse_run, sample
from .lining.model import angle_breaks, validate
from .lining.report import build_report

LIB_PATH = Path(__file__).parent / "data" / "brick_library_mock.json"


class LiningInputError(ValueError):
    def __init__(self, errors):
        super().__init__("; ".join(errors))
        self.errors = errors


def load_library(payload_lib=None):
    if payload_lib and payload_lib.get("bricks"):
        return payload_lib
    return json.loads(LIB_PATH.read_text(encoding="utf-8"))


def add_entered_bricks(lib, payload, st):
    """Brick shapes entered in the UI steps (zone brick pairs, mouth pairs, bottom ring bands, permanent
    brick) are added to the library so drawings, BOM and checks use exactly those dimensions.
    An entered mark that also exists in the library with other dimensions replaces it (UI wins)."""
    bricks = {b["mark"]: b for b in lib.get("bricks", [])}
    notes, added = [], {}

    def put(b):
        old = bricks.get(b["mark"])
        if old and old.get("dimensions_mm") != b["dimensions_mm"] and b["mark"] not in added:
            notes.append(f"{b['mark']}: dimensions entered in the UI used instead of the brick library")
        if b["mark"] in added and added[b["mark"]]["dimensions_mm"] != b["dimensions_mm"]:
            notes.append(f"{b['mark']}: entered twice with different dimensions - check the brick pairs")
            return
        bricks[b["mark"]] = b
        added[b["mark"]] = b

    design = payload.get("design") or {}
    for band in design.get("bands", []):
        kind = band.get("kind") or "wall"
        ch = float(band.get("course_height") or (st.get("mouth_course_height") if kind == "mouth" else st.get("course_height")) or 100)
        brick_h = float(band.get("brick_height") or ch)
        for sec in band.get("sectors", []):
            grade = (sec.get("layers") or [{}])[-1].get("grade", "")
            for q in sec.get("pairs") or []:
                try:
                    L = float(q["L"])
                    dims = [(str(q["A"]), float(q["aCold"]), float(q["aHot"])), (str(q["B"]), float(q["bCold"]), float(q["bHot"]))]
                except (KeyError, TypeError, ValueError):
                    continue
                fam = q.get("family") or "standard_trapezoidal"
                for mark, cold, hot in dims:
                    if min(L, cold, hot) <= 0:
                        continue
                    put({"mark": mark, "shape_family": fam, "zone": f"{band['name']} / {sec.get('name', '')}",
                         "quality_grade_internal": grade, "source": "entered in UI",
                         "dimensions_mm": {"length": L, "width_top": cold, "width_bottom": hot, "thickness": brick_h}})
    bot = payload.get("bottom") or {}
    rbs = bot.get("ring_bands") or []
    if rbs and all(all(str(rb.get(k, "")).strip() not in ("", "0") for k in ("aHot", "aCold", "bHot", "bCold")) for rb in rbs):
        w, L = float(bot.get("ring_width") or 0), float(bot.get("brick_length") or 0)
        for rb in rbs:
            for suf, inner, outer in (("A", rb["aHot"], rb["aCold"]), ("B", rb["bHot"], rb["bCold"])):
                put({"mark": f"{rb['mark']}{suf}", "shape_family": "bottom_dish_wedge", "zone": "Bottom",
                     "quality_grade_internal": bot.get("grade", ""), "source": "entered in UI",
                     "dimensions_mm": {"overall_length": L, "face_width_top": w, "face_width_bottom": w,
                                       "A": float(outer), "B": float(inner), "C": float(outer), "D": float(inner)}})
    pb = payload.get("permanent_brick") or {}
    if pb.get("width") and pb.get("height"):
        ts = set()
        mono = payload.get("monolithic") or {}
        for band in design.get("bands", []):
            for sec in band.get("sectors", []):
                lay = sec.get("layers") or []
                if len(lay) > 1 and lay[0].get("grade") not in mono:      # castable / filling is not bricks
                    ts.add(float(lay[0]["thickness"]))
        for t in sorted(ts):
            if t > 0:
                put({"mark": f"PL-{float(pb['width']):g}x{float(pb['height']):g}x{t:g}", "shape_family": "safety_lining_std",
                     "zone": "Permanent lining", "quality_grade_internal": "", "source": "entered in UI",
                     "dimensions_mm": {"length": t, "width": float(pb["width"]), "thickness": float(pb["height"])}})
        # the library's own permanent bricks are not used when the permanent brick is entered
        for m in [m for m, b in bricks.items() if b.get("shape_family") == "safety_lining_std" and m not in added]:
            del bricks[m]
    out = dict(lib)
    out["bricks"] = list(bricks.values())
    return out, notes


def _num_or_none(v):
    if v in (None, ""):
        return None
    try:
        f = float(v)
        return None if math.isnan(f) else f
    except (TypeError, ValueError):
        return None


def _nozzles(raw):
    out = []
    for n in raw or []:
        t = (n.get("type") or "radial").lower()
        d = {"label": str(n.get("label") or t), "type": t}
        for k in ("z", "angle", "inclination", "diameter", "count", "pcd", "start_angle", "size_a", "size_b", "height"):
            d[k] = _num_or_none(n.get(k))
        if t == "stirring":
            if not d["count"] or not d["pcd"]:
                continue
            d["count"] = int(d["count"])
        elif d["z"] is None or d["angle"] is None:
            continue
        out.append(d)
    return out


def _aci_colors(colors):
    """UI colours (#rrggbb) -> nearest AutoCAD colour index, so the drawing matches the wizard."""
    from ezdxf.colors import DXF_DEFAULT_COLORS, int2rgb
    out = {}
    for g, c in colors.items():
        if isinstance(c, str) and re.fullmatch(r"#?[0-9a-fA-F]{6}", c.strip()):
            h = c.strip().lstrip("#")
            rgb = tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))
            best = min(range(1, 256), key=lambda i: sum((a - b) ** 2 for a, b in zip(int2rgb(DXF_DEFAULT_COLORS[i]), rgb)))
            out[g] = best
        elif isinstance(c, (int, float)) or (isinstance(c, str) and c.strip()):
            out[g] = c
    return out


def volumes(sheet):
    """Structured version of report.build_report (same maths)."""
    from refractory_core import get_core
    f = get_core().formulas
    k = 1e-9 if sheet.vessel.get("units", "mm") == "mm" else 1.0
    L = sheet.L
    br = angle_breaks(sheet.design)
    cavity = 0.0
    for a, b in zip(br, br[1:]):
        paths, _ = sheet.boundary_paths((a + b) / 2)
        hot = paths[L].all()
        pts = sample(hot) + [(0.0, hot[-1][1]), (0.0, hot[0][1])]
        cavity += f.evaluate("revolved_volume", {
            "area_moment": region_area_moment(pts)[1],
            "angle_from": a, "angle_to": b, "unit_factor": 1.0,
        })
    rows, by_grade, by_layer = [], {}, {}
    for bi, band in enumerate(sheet.bands):
        i0, i1 = sheet.band_idx[bi]
        for sec in band["sectors"]:
            paths, _ = sheet.boundary_paths((sec["from"] + sec["to"]) / 2)
            for j in range(1, L + 1):
                pts = sample(paths[j - 1].run(i0, i1) + reverse_run(paths[j].run(i0, i1)))
                vol = f.evaluate("revolved_volume", {
                    "area_moment": region_area_moment(pts)[1],
                    "angle_from": sec["from"], "angle_to": sec["to"], "unit_factor": k,
                })
                l = sec["layers"][j - 1]
                rows.append({"band": band["name"], "sector": f"{sec['from']:g}-{sec['to']:g}", "zone": sec.get("name", ""),
                             "layer": sheet.layers[j - 1], "grade": l["grade"], "thickness": l["thickness"],
                             "volume_m3": round(vol, 3)})
                by_grade[l["grade"]] = by_grade.get(l["grade"], 0.0) + vol
                by_layer[sheet.layers[j - 1]] = by_layer.get(sheet.layers[j - 1], 0.0) + vol
    hs = sheet.design.get("heat_size")
    cav_m3 = cavity * k
    return {"rows": rows, "by_grade": {g: round(v, 3) for g, v in by_grade.items()},
            "by_layer": {g: round(v, 3) for g, v in by_layer.items()},
            "total_lining_m3": round(sum(by_layer.values()), 3), "inside_volume_m3": round(cav_m3, 3),
            "specific_volume_m3_per_t": round(f.evaluate("specific_volume", {"cavity_m3": cav_m3, "heat_size": float(hs)}), 4) if hs else None}


def _write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def _write_xlsx(path, job, head, rows, bom):
    """BOM workbook (sheet BOM + sheet By grade). Returns False if openpyxl is missing."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.append([f"Bill of materials - {job}  (generated from the lining design; cut allowance {bom['cut_allowance_pct']:g} %)"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(head)
    for c in ws[3]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E6F4EC")
    for r in rows:
        ws.append(r)
    ws.append([])
    ws.append(["TOTAL", "", "", "", "", "", "", "", "", bom["pieces"], "", "", "", bom["total_t"], ""])
    ws[ws.max_row][0].font = Font(bold=True)
    for col, w in zip("ABCDEFGHIJKLMNO", (5, 24, 30, 18, 22, 20, 26, 8, 10, 10, 11, 10, 10, 10, 60)):
        ws.column_dimensions[col].width = w
    g = wb.create_sheet("By grade")
    g.append(["Grade", "Tonnes"])
    for k, v in bom["by_grade"].items():
        g.append([k, v])
    g.append(["TOTAL", bom["total_t"]])
    wb.save(path)
    return True


def generate(payload: dict, out_dir: Path, url_prefix: str = "", progress=None) -> dict:
    """progress(stage_text, percent) is called as the work advances (optional)."""
    step = progress or (lambda *_: None)
    step("Checking the input values", 2)
    out_dir.mkdir(parents=True, exist_ok=True)
    lib = load_library(payload.get("brick_library"))
    from refractory_core import get_core
    f = get_core().formulas
    st = {"course_height": f.evaluate("d_course_height", {}),
          "mouth_course_height": f.evaluate("d_mouth_course_height", {}),
          "joint": f.evaluate("d_joint", {}),
          "bottom_joint": f.evaluate("d_bottom_joint", {}),
          "min_cut": f.evaluate("d_min_cut", {}),
          "cut_allowance": f.evaluate("d_cut_allowance", {}),
          **{k: v for k, v in (payload.get("settings") or {}).items() if v not in (None, "")}}
    for k in list(st):
        try:
            st[k] = float(st[k])
        except (TypeError, ValueError):
            raise LiningInputError([f"setting '{k}': '{st[k]}' is not a number"])
    lib, lib_notes = add_entered_bricks(lib, payload, st)
    nozzles = _nozzles(payload.get("nozzles"))
    job = str(payload.get("job") or "job")

    payload = dict(payload)
    if (payload.get("design") or {}).get("grade_colors"):
        payload["design"] = {**payload["design"], "grade_colors": _aci_colors(payload["design"]["grade_colors"])}
    try:
        vessel, design, _ = build_engine_inputs(payload)
    except (KeyError, ValueError) as e:
        raise LiningInputError([f"{e}"])
    vessel["job"] = job
    vessel.setdefault("units", "mm")
    design["course_height"] = st["course_height"]
    err, warn = validate(vessel, design)
    if err:
        raise LiningInputError(err)

    try:
        sheet = LiningSheet(vessel, design, nozzles)
        views = design["views"]
        if not views.get("plan_heights"):
            hs = [round(sum(sheet.band_z(b)) / 2, 1) for b in range(len(sheet.bands))]
            hs += [n["z"] for n in nozzles if n["type"] != "stirring" and sheet.zmin < n["z"] < sheet.zmax]
            views["plan_heights"] = sorted(set(hs))
        step("Drawing the lining GA (sections, plans, top view, unrolled map, nozzles)", 8)
        sheet.build()
    except ValueError as e:
        raise LiningInputError([f"Could not build the lining drawing: {e}"])

    files = {}
    ga = out_dir / "lining_ga.dxf"
    sheet.doc.saveas(ga)
    files["lining_ga.dxf"] = ga
    step("Rendering the Lining GA preview (PNG / PDF)", 15)
    for p in save_previews(sheet.msp, out_dir, "lining_ga"):
        files[p.name] = p

    report = build_report(sheet)
    vols = volumes(sheet)
    step("Computing course schedule and bottom rings from the brick library", 35)
    courses, course_flags, cone_gaps = course_schedule(sheet, lib, st)
    rings = bottom_rings(sheet, lib, st, nozzles, payload.get("bottom") or {})
    qty = quantities(courses, rings)
    bom = build_bom(sheet, lib, courses, rings, vols, nozzles, st, payload.get("grade_density") or {}, cone_gaps,
                    payload.get("monolithic") or {}, payload.get("extra_items") or [], payload.get("junctions") or [])
    bom["flags"] = sorted(set(bom["flags"]) | set(lib_notes))

    step("Checking the drawings against the input values", 38)
    checks = run_checks({**payload, "brick_library": lib}, vessel, sheet, courses, rings, vols, bom, st, nozzles)
    mcol = assign_mark_colors(courses, rings)
    minfo = mark_info(lib, courses, rings, qty, mcol)
    qty_sheet = dict(qty)                                   # + permanent-lining bricks counted in the BOM
    for o in bom["lines"]:
        if o["group"] == "Permanent lining" and o.get("qty"):
            qty_sheet[o["mark"]] = qty_sheet.get(o["mark"], 0) + o["qty"]
    bs = BrickSheet(lib, qty_sheet, job, mcol, minfo)
    step("Drawing the brick shapes sheet", 40)
    bs.build()
    p = out_dir / "brick_shapes.dxf"
    bs.doc.saveas(p)
    files[p.name] = p
    step("Rendering the Brick shapes preview (PNG / PDF)", 45)
    for q in save_previews(bs.msp, out_dir, "brick_shapes", size=(36, 30)):
        files[q.name] = q

    ls = LayerSheet(sheet, courses, rings, nozzles, job, mcol, minfo)
    step("Drawing the layers & brick layout sheet", 65)
    ls.build()
    p = out_dir / "lining_layers.dxf"
    ls.doc.saveas(p)
    files[p.name] = p
    step("Rendering the Layers & brick layout preview (PNG / PDF)", 70)
    for q in save_previews(ls.msp, out_dir, "lining_layers", size=(40, 24)):
        files[q.name] = q

    # tables
    step("Writing tables (CSV) and summary", 95)
    _write_csv(out_dir / "course_schedule.csv",
               ["course", "z0", "z1", "height", "band", "sector", "zone", "grade", "brick_length", "r_cold", "r_hot",
                "pair", "mark_A", "nA", "mark_B", "nB", "cut_pieces", "cut_mm", "gap_mm", "flag"],
               [[r["course"], r["z0"], r["z1"], r["h"], r["band"], r["sector"], r.get("zone", ""), r["grade"], r["L"],
                 r.get("r_cold"), r.get("r_hot"), r.get("pair") or "", r.get("A") or "", r["nA"], r.get("B") or "",
                 r["nB"], r["nCut"], r["cut"], r.get("gap"), r.get("flag") or ""] for r in courses])
    files["course_schedule.csv"] = out_dir / "course_schedule.csv"
    if rings:
        _write_csv(out_dir / "bottom_rings.csv", ["ring", "r_in", "r_out", "marks", "cut_pieces", "cut_mm", "note", "flag"],
                   ([[0, 0, rings["key"]["r"], f"{rings['key']['mark']} x 1", 0, 0, "", ""]] if rings.get("key") else []) +
                   [[r["ring"], r["r_in"], r["r_out"], " + ".join(f"{m} x {q}" for m, q in r["marks"].items()), r["nCut"],
                     r["cut"], r.get("note") or "", r.get("flag") or ""] for r in rings["rings"]])
        files["bottom_rings.csv"] = out_dir / "bottom_rings.csv"
    bom_head = ["item", "group", "zone", "grade", "mark", "shape", "size_mm", "qty", "allowance_pct", "order_qty",
                "unit_vol_dm3", "density_g_cm3", "unit_wt_kg", "total_t", "note"]
    bom_rows = [[o["item"], o["group"], o["zone"], o["grade"], o["mark"], o["family"], o["size"], o["qty"],
                 o["allow_pct"], o["order_qty"], o["unit_vol_dm3"], o["density"], o["unit_wt_kg"], o["total_t"],
                 (o.get("note") or "") + (f" [volume {o['volume_m3']} m3]" if o.get("volume_m3") else "")]
                for o in bom["lines"]]
    _write_csv(out_dir / "bom.csv", bom_head, bom_rows + [[], ["TOTAL", "", "", "", "", "", "", "", "",
               bom["pieces"], "", "", "", bom["total_t"], ""]])
    files["bom.csv"] = out_dir / "bom.csv"
    if _write_xlsx(out_dir / "bom.xlsx", job, bom_head, bom_rows, bom):
        files["bom.xlsx"] = out_dir / "bom.xlsx"
    _write_csv(out_dir / "brick_quantities.csv", ["mark", "pieces"], sorted(qty.items()))
    files["brick_quantities.csv"] = out_dir / "brick_quantities.csv"

    flags = list(course_flags) + list((rings or {}).get("flags", []))
    lib_gaps = [f"{g.get('issue')}: {g.get('status')}" for g in lib.get("_known_gaps", [])]
    summary = [report, "", "BRICK QUANTITIES (pieces, before allowances)"]
    summary += [f"  {m:<24}{n:>8}" for m, n in sorted(qty.items())]
    summary += ["", "ACCURACY CHECKS (drawing measured again and compared with the inputs)"]
    summary += [f"  [{'PASS' if c['ok'] else 'FAIL'}] {c['check']}: {c['value']} (expected {c['expected']}) {c['detail']}"
                for c in checks]
    summary += ["", "FLAGS"] + [f"  - {f}" for f in flags] + ["", "BRICK LIBRARY GAPS"] + [f"  - {g}" for g in lib_gaps]
    (out_dir / "summary.txt").write_text("\n".join(summary) + "\n", encoding="utf-8")
    files["summary.txt"] = out_dir / "summary.txt"
    (out_dir / "vessel.json").write_text(json.dumps(vessel, indent=2), encoding="utf-8")
    (out_dir / "design.json").write_text(json.dumps(design, indent=2), encoding="utf-8")
    (out_dir / "input_payload.json").write_text(json.dumps({k: v for k, v in payload.items() if k != "brick_library"},
                                                           indent=2, default=str), encoding="utf-8")
    files["vessel.json"] = out_dir / "vessel.json"
    files["design.json"] = out_dir / "design.json"

    # one ZIP with every output, each file named "<project>_<file>"
    import zipfile
    prefix = safe_name(str(payload.get("file_prefix") or job))
    pkg = out_dir / "package.zip"
    with zipfile.ZipFile(pkg, "w", zipfile.ZIP_DEFLATED) as z:
        for name, path in files.items():
            z.write(path, arcname=f"{prefix}_{name}")
    files["package.zip"] = pkg

    n_wall = len({r["course"] for r in courses})
    return {
        "job": job,
        "file_prefix": safe_name(str(payload.get("file_prefix") or job)),
        "generated_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "files": {name: f"{url_prefix}/{name}" for name in files},
        "report": report,
        "volumes": vols,
        "courses": courses,
        "n_courses": n_wall,
        "bottom_rings": rings,
        "quantities": qty,
        "bom": bom,
        "checks": checks,
        "mark_colors": {m: aci_hex(c) for m, c in mcol.items()},
        "mark_info": minfo,
        "flags": flags,
        "library_gaps": lib_gaps,
        "warnings": warn,
        "engine_vessel": vessel,
        "bands": [{"name": b["name"], "kind": band_kind(sheet, i), "z": [round(x, 1) for x in sheet.band_z(i)]}
                  for i, b in enumerate(sheet.bands)],
    }


def safe_name(s: str) -> str:
    return re.sub(r"[^\w\-]+", "_", s or "job").strip("_") or "job"
